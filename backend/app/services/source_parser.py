from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import date, datetime
from itertools import product
from typing import Any

import xlrd
from openpyxl import load_workbook

from app.core.errors import DomainError
from app.core.sensitive import classify_field


def checksum(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _value_type(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, dict | list):
        return "object" if isinstance(value, dict) else "array"
    return "string"


def _schema(rows: list[dict[str, Any]]) -> dict[str, Any]:
    field_names: list[str] = []
    values: dict[str, Any] = {}
    for row in rows:
        for name, value in row.items():
            if name not in field_names:
                field_names.append(name)
            if name not in values and value is not None:
                values[name] = value
    columns = []
    for name in field_names:
        decision = classify_field(name, _value_type(values.get(name, "")))
        columns.append(
            {
                "name": name,
                "type": _value_type(values.get(name)),
                "classification": decision.classification,
                "action": decision.action,
                "confidence": decision.confidence,
                "reason": decision.reason,
            }
        )
    return {"row_count": len(rows), "columns": columns}


def _validate_limits(rows: list[dict[str, Any]], max_rows: int, max_columns: int) -> None:
    if len(rows) > max_rows:
        raise DomainError("O arquivo excede o limite de linhas.", "source_limit_exceeded", 413)
    if len({key for row in rows for key in row}) > max_columns:
        raise DomainError("O arquivo excede o limite de colunas.", "source_limit_exceeded", 413)


def _validate_json_records(rows: list[dict[str, Any]], configuration: dict[str, Any]) -> None:
    if not rows:
        return
    allow_missing = bool(configuration.get("allow_missing_fields", True))
    if not allow_missing:
        expected_fields = set(rows[0])
        if any(set(row) != expected_fields for row in rows[1:]):
            raise DomainError(
                "O JSON possui campos ausentes ou extras entre registros.",
                "invalid_source_file",
            )
    field_types: dict[str, set[str]] = {}
    for row in rows:
        for field, value in row.items():
            field_types.setdefault(field, set()).add(_value_type(value))
    conflicts = {
        field: types
        for field, types in field_types.items()
        if len(types - {"null"}) > 1
        and not (types - {"null"}).issubset({"integer", "number"})
    }
    if not conflicts:
        return
    strategy = str(configuration.get("type_conflict_strategy", "FAIL_VALIDATION"))
    if strategy != "COERCE_STRING":
        fields = ", ".join(sorted(conflicts))
        raise DomainError(
            f"Conflito de tipo nos campos: {fields}.", "invalid_source_file"
        )
    for row in rows:
        for field in conflicts:
            if row.get(field) is not None:
                row[field] = str(row[field])


def parse_csv(
    content: bytes, configuration: dict[str, Any], max_rows: int, max_columns: int
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    encoding = str(configuration.get("encoding", "utf-8-sig"))
    delimiter = str(configuration.get("delimiter", ","))
    has_header = bool(configuration.get("has_header", True))
    header_row = int(configuration.get("header_row", 0))
    try:
        text = content.decode(encoding)
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except (UnicodeDecodeError, csv.Error, ValueError) as exc:
        raise DomainError(
            "CSV inválido ou incompatível com a configuração.", "invalid_source_file"
        ) from exc
    if not rows or header_row >= len(rows):
        raise DomainError("CSV sem linhas suficientes para o cabeçalho.", "invalid_source_file")
    raw_headers = (
        rows[header_row] if has_header else [f"coluna_{index + 1}" for index in range(len(rows[0]))]
    )
    headers = [header.strip() or f"coluna_{index + 1}" for index, header in enumerate(raw_headers)]
    if len(headers) != len(set(headers)):
        raise DomainError("O CSV possui nomes de coluna duplicados.", "invalid_source_file")
    data_rows = rows[header_row + 1 :] if has_header else rows
    parsed = [dict(zip(headers, row, strict=False)) for row in data_rows]
    _validate_limits(parsed, max_rows, max_columns)
    return parsed, _schema(parsed)


def _worksheet_rows(
    values: list[tuple[Any, ...]], header_row: int, max_rows: int
) -> list[dict[str, Any]]:
    if not values or header_row >= len(values):
        raise DomainError(
            "Planilha sem linhas suficientes para o cabeçalho.", "invalid_source_file"
        )
    headers = [str(value).strip() if value is not None else "" for value in values[header_row]]
    headers = [header or f"coluna_{index + 1}" for index, header in enumerate(headers)]
    if len(headers) != len(set(headers)):
        raise DomainError("A planilha possui nomes de coluna duplicados.", "invalid_source_file")
    return [
        dict(zip(headers, row, strict=False))
        for row in values[header_row + 1 : max_rows + header_row + 1]
    ]


def _apply_excel_policies(
    rows: list[dict[str, Any]],
    configuration: dict[str, Any],
    formula_cells: set[tuple[int, int]],
    headers: list[str],
    header_row: int,
) -> None:
    """Apply explicit, deterministic policies for ambiguous spreadsheet cells."""

    formula_strategy = str(configuration.get("formula_strategy", "USE_CACHED_VALUE")).upper()
    empty_strategy = str(configuration.get("empty_cell_strategy", "NULL")).upper()
    consistency_strategy = str(
        configuration.get("column_consistency_strategy", "PAD_MISSING_WITH_NULL")
    ).upper()
    if formula_strategy not in {"USE_CACHED_VALUE", "REJECT"}:
        raise DomainError("Estratégia de fórmula Excel inválida.", "invalid_import_config")
    if empty_strategy not in {"NULL", "EMPTY_STRING", "REJECT"}:
        raise DomainError("Estratégia de célula vazia inválida.", "invalid_import_config")
    if consistency_strategy not in {"PAD_MISSING_WITH_NULL", "REJECT"}:
        raise DomainError("Estratégia de colunas inconsistentes inválida.", "invalid_import_config")

    for row_number, row in enumerate(rows, start=header_row + 2):
        if consistency_strategy == "REJECT" and len(row) != len(headers):
            raise DomainError(
                f"A linha {row_number} possui quantidade de colunas inconsistente.",
                "invalid_source_file",
            )
        for column_index, field_name in enumerate(headers):
            value = row.get(field_name)
            coordinate = (row_number, column_index + 1)
            if coordinate in formula_cells and formula_strategy == "REJECT":
                raise DomainError(
                    f"A célula com fórmula {field_name} na linha {row_number} não é permitida.",
                    "invalid_source_file",
                )
            if value is None:
                if coordinate in formula_cells:
                    continue
                if empty_strategy == "REJECT":
                    raise DomainError(
                        f"A célula vazia {field_name} na linha {row_number} não é permitida.",
                        "invalid_source_file",
                    )
                if empty_strategy == "EMPTY_STRING":
                    row[field_name] = ""

    date_columns = {str(item) for item in configuration.get("date_columns", [])}
    invalid_date_strategy = str(
        configuration.get("invalid_date_strategy", "KEEP_TEXT")
    ).upper()
    if invalid_date_strategy not in {"KEEP_TEXT", "COERCE_NULL", "REJECT"}:
        raise DomainError("Estratégia de datas inválida.", "invalid_import_config")
    date_format = str(configuration.get("date_format", "%Y-%m-%d"))
    for row_number, row in enumerate(rows, start=header_row + 2):
        for field_name in date_columns:
            value = row.get(field_name)
            if value in (None, "") or isinstance(value, date | datetime):
                continue
            try:
                datetime.strptime(str(value), date_format)
            except ValueError as exc:
                if invalid_date_strategy == "REJECT":
                    raise DomainError(
                        f"A data da coluna {field_name} na linha {row_number} é inválida.",
                        "invalid_source_file",
                    ) from exc
                if invalid_date_strategy == "COERCE_NULL":
                    row[field_name] = None


def parse_excel(
    content: bytes, configuration: dict[str, Any], max_rows: int, max_columns: int
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    extension = str(configuration.get("extension", ".xlsx")).lower()
    sheet_name = configuration.get("sheet_name")
    header_row = int(configuration.get("header_row", 0))
    try:
        if extension == ".xls":
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            sheet = workbook.sheet_by_name(sheet_name) if sheet_name else workbook.sheet_by_index(0)
            values = [
                tuple(sheet.row_values(index))
                for index in range(min(sheet.nrows, max_rows + header_row + 1))
            ]
            formula_cells: set[tuple[int, int]] = set()
        else:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
            sheet_protection = getattr(getattr(sheet, "protection", None), "sheet", False)
            if bool(sheet_protection) and str(
                configuration.get("protected_sheet_strategy", "ALLOW_READ_ONLY")
            ).upper() == "REJECT":
                raise DomainError(
                    "A planilha protegida não pode ser importada com esta política.",
                    "invalid_source_file",
                )
            values = [
                tuple(row)
                for row in sheet.iter_rows(max_row=max_rows + header_row + 1, values_only=True)
            ]
            workbook.close()
            formula_workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            formula_sheet = (
                formula_workbook[sheet_name]
                if sheet_name
                else formula_workbook.worksheets[0]
            )
            formula_cells = {
                (row_index, cell.column)
                for row_index, row in enumerate(
                    formula_sheet.iter_rows(max_row=max_rows + header_row + 1), start=1
                )
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            }
            formula_workbook.close()
        headers = [
            str(value).strip() if value is not None else ""
            for value in values[header_row]
        ]
        headers = [header or f"coluna_{index + 1}" for index, header in enumerate(headers)]
        parsed = _worksheet_rows(values, header_row, max_rows)
        _apply_excel_policies(parsed, configuration, formula_cells, headers, header_row)
    except DomainError:
        raise
    except (ValueError, KeyError, IndexError, TypeError, xlrd.XLRDError) as exc:
        raise DomainError(
            "Excel inválido ou incompatível com a configuração.", "invalid_source_file"
        ) from exc
    _validate_limits(parsed, max_rows, max_columns)
    return parsed, _schema(parsed)


def _json_depth(value: Any, current: int = 0) -> int:
    if isinstance(value, dict):
        return max([current, *(_json_depth(item, current + 1) for item in value.values())])
    if isinstance(value, list):
        return max([current, *(_json_depth(item, current + 1) for item in value)])
    return current


def _path_value(payload: Any, root_path: str) -> Any:
    if not root_path or root_path == "$":
        return payload
    current = payload
    for part in root_path.removeprefix("$.").split("."):
        if not isinstance(current, dict) or part not in current:
            raise DomainError("root_path não encontrado no JSON.", "invalid_source_file")
        current = current[part]
    return current


def _expand_json(value: Any, prefix: str, separate_arrays: bool) -> list[dict[str, Any]]:
    if isinstance(value, dict):
        rows: list[dict[str, Any]] = [{}]
        for key, item in value.items():
            child_rows = _expand_json(item, f"{prefix}.{key}".strip("."), separate_arrays)
            rows = [{**left, **right} for left, right in product(rows, child_rows)]
        return rows
    if isinstance(value, list) and separate_arrays:
        expanded_rows: list[dict[str, Any]] = []
        for item in value:
            expanded_rows.extend(_expand_json(item, prefix, separate_arrays))
        return expanded_rows or [{prefix: None}]
    return [{prefix: value}]


def parse_json(
    content: bytes, configuration: dict[str, Any], max_rows: int, max_columns: int, max_depth: int
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        payload = json.loads(content.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise DomainError("JSON inválido.", "invalid_source_file") from exc
    if _json_depth(payload) > max_depth:
        raise DomainError("JSON excede a profundidade máxima.", "source_limit_exceeded", 413)
    records = _path_value(payload, str(configuration.get("root_path", "$")))
    if isinstance(records, dict):
        records = [records]
    if not isinstance(records, list):
        raise DomainError(
            "A raiz de registros do JSON deve ser uma lista ou objeto.", "invalid_source_file"
        )
    separate_arrays = configuration.get("nested_array_strategy", "REJECT") == "SEPARATE_RECORDS"
    if any(_contains_nested_array(item) and not separate_arrays for item in records):
        raise DomainError("Arrays aninhados exigem estratégia explícita.", "invalid_source_file")
    parsed: list[dict[str, Any]] = []
    for record in records:
        if not isinstance(record, dict):
            raise DomainError("Cada registro JSON deve ser um objeto.", "invalid_source_file")
        parsed.extend(_expand_json(record, "", separate_arrays))
    _validate_json_records(parsed, configuration)
    _validate_limits(parsed, max_rows, max_columns)
    return parsed, _schema(parsed)


def _contains_nested_array(value: Any) -> bool:
    if isinstance(value, list):
        return True
    if isinstance(value, dict):
        return any(_contains_nested_array(item) for item in value.values())
    return False


def parse_file(
    source_type: str,
    content: bytes,
    configuration: dict[str, Any],
    max_rows: int,
    max_columns: int,
    max_json_depth: int,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if source_type == "ARQUIVO_CSV":
        return parse_csv(content, configuration, max_rows, max_columns)
    if source_type == "ARQUIVO_EXCEL":
        return parse_excel(content, configuration, max_rows, max_columns)
    if source_type == "ARQUIVO_JSON":
        return parse_json(content, configuration, max_rows, max_columns, max_json_depth)
    raise DomainError("Tipo de arquivo não suportado.", "unsupported_source_type")
