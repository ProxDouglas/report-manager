import io
import json

import pytest
from openpyxl import Workbook

from app.core.errors import DomainError
from app.core.security import hash_password, verify_password
from app.core.sensitive import classify_field, mask_sensitive_text, scrub_mapping
from app.db.enums import DataClassification, SensitivityAction
from app.services.connectors import validate_public_host, validate_read_only_query
from app.services.source_parser import parse_csv, parse_excel, parse_json


def test_password_hash_is_not_reversible() -> None:
    password_hash = hash_password("StrongLocal#12345")
    assert password_hash != "StrongLocal#12345"
    assert verify_password(password_hash, "StrongLocal#12345")
    assert not verify_password(password_hash, "wrong-password")


def test_sensitive_values_are_masked_and_secrets_are_redacted() -> None:
    decision = classify_field("password_hash")
    assert decision.classification is DataClassification.SECRET
    assert decision.action is SensitivityAction.BLOCK
    values = scrub_mapping({"cpf": "12345678909", "password": "do-not-log"})
    assert values["cpf"] == "***.***.***-09"
    assert values["password"] == "[REDACTED]"
    assert "12345678909" not in mask_sensitive_text("CPF 12345678909")


def test_csv_json_and_excel_parsers_apply_schema() -> None:
    rows, schema = parse_csv(b"nome,valor\nA,10\nB,20\n", {}, 10, 10)
    assert rows[0]["nome"] == "A"
    assert schema["row_count"] == 2

    json_rows, json_schema = parse_json(
        json.dumps({"clientes": [{"id": 1, "dados": {"cidade": "SP"}}]}).encode(),
        {"root_path": "$.clientes"},
        10,
        10,
        5,
    )
    assert json_rows[0]["dados.cidade"] == "SP"
    assert json_schema["columns"][1]["name"] == "dados.cidade"

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["nome", "valor"])
    sheet.append(["A", 10])
    content = io.BytesIO()
    workbook.save(content)
    excel_rows, _ = parse_excel(content.getvalue(), {"extension": ".xlsx"}, 10, 10)
    assert excel_rows == [{"nome": "A", "valor": 10}]


def test_json_depth_and_sql_write_are_rejected() -> None:
    with pytest.raises(DomainError):
        parse_json(b'{"a":{"b":{"c":1}}}', {}, 10, 10, 1)
    with pytest.raises(DomainError):
        validate_read_only_query("UPDATE clientes SET nome = 'x'")
    with pytest.raises(DomainError):
        validate_public_host("10.0.0.1")


def test_json_type_conflict_requires_explicit_strategy() -> None:
    payload = b'{"items":[{"value":1},{"value":"texto"}]}'
    with pytest.raises(DomainError):
        parse_json(payload, {"root_path": "$.items"}, 10, 10, 5)
    rows, _ = parse_json(
        payload,
        {"root_path": "$.items", "type_conflict_strategy": "COERCE_STRING"},
        10,
        10,
        5,
    )
    assert rows == [{"value": "1"}, {"value": "texto"}]


def test_json_nested_arrays_and_missing_fields_follow_explicit_policy() -> None:
    payload = b'{"items":[{"id":1,"tags":["a","b"]},{"id":2}]}'
    with pytest.raises(DomainError):
        parse_json(payload, {"root_path": "$.items"}, 10, 10, 5)
    rows, _ = parse_json(
        payload,
        {
            "root_path": "$.items",
            "nested_array_strategy": "SEPARATE_RECORDS",
            "allow_missing_fields": True,
        },
        10,
        10,
        5,
    )
    assert len(rows) == 3
    assert rows[-1].get("tags") is None


def test_excel_policies_define_empty_cells_formulas_and_invalid_dates() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["data", "valor", "formula"])
    sheet.append(["2026-01-01", None, "=1+1"])
    content = io.BytesIO()
    workbook.save(content)

    rows, _ = parse_excel(
        content.getvalue(),
        {
            "extension": ".xlsx",
            "date_columns": ["data"],
            "empty_cell_strategy": "EMPTY_STRING",
        },
        10,
        10,
    )
    assert rows[0]["valor"] == ""
    assert rows[0]["formula"] is None
    with pytest.raises(DomainError):
        parse_excel(
            content.getvalue(),
            {"extension": ".xlsx", "formula_strategy": "REJECT"},
            10,
            10,
        )
    workbook = Workbook()
    workbook.active.append(["data"])
    workbook.active.append(["not-a-date"])
    invalid_content = io.BytesIO()
    workbook.save(invalid_content)
    with pytest.raises(DomainError):
        parse_excel(
            invalid_content.getvalue(),
            {"extension": ".xlsx", "date_columns": ["data"], "invalid_date_strategy": "REJECT"},
            10,
            10,
        )
